#!/usr/bin/env python3
"""
Build the GCC Atlas presentation deck (PowerPoint .pptx).

Run with:   python tools/build-deck.py
Output:     deck/GCC-Atlas-Flexiple-Assignment.pptx

Design follows the same restraint as the app: one accent colour, a warm
off-white background, serif display type, generous margins, no decoration.
Screenshots are cropped from the real application, not mocked up.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHOTS = os.path.join(ROOT, "screenshots")
DECK_DIR = os.path.join(ROOT, "deck")
CROPS = os.path.join(SHOTS, "crops")
os.makedirs(DECK_DIR, exist_ok=True)
os.makedirs(CROPS, exist_ok=True)

# ----------------------------------------------------------------- palette --
INK        = RGBColor(0x1A, 0x1A, 0x1A)
INK2       = RGBColor(0x4A, 0x48, 0x45)
INK3       = RGBColor(0x7A, 0x77, 0x72)
ACCENT     = RGBColor(0x1F, 0x3A, 0x5F)
ACCENT_BG  = RGBColor(0xEA, 0xEE, 0xF4)
BG         = RGBColor(0xFB, 0xFA, 0xF8)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
LINE       = RGBColor(0xE2, 0xDE, 0xD7)
GREEN      = RGBColor(0x1D, 0x5C, 0x43)
GREEN_BG   = RGBColor(0xE8, 0xF1, 0xED)
AMBER      = RGBColor(0x7A, 0x56, 0x18)
AMBER_BG   = RGBColor(0xF7, 0xF0, 0xE2)
RED        = RGBColor(0x7D, 0x3B, 0x36)
RED_BG     = RGBColor(0xF8, 0xEC, 0xEB)
PURPLE     = RGBColor(0x6B, 0x4A, 0x7A)
PURPLE_BG  = RGBColor(0xF4, 0xEF, 0xF6)

SERIF = "Georgia"
SANS = "Calibri"

W, H = Inches(13.333), Inches(7.5)
M = Inches(0.78)              # page margin
CONTENT_W = W - 2 * M

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]

slide_no = {"n": 0}


# ------------------------------------------------------------------ helpers --
def add_slide(bg=BG):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(1, 0, 0, W, H)      # 1 = rectangle
    r.fill.solid()
    r.fill.fore_color.rgb = bg
    r.line.fill.background()
    r.shadow.inherit = False
    return s


def textbox(s, x, y, w, h, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    tf.paragraphs[0].alignment = align
    return tf


def para(tf, text, size=14, color=INK2, bold=False, font=SANS, space_after=6,
         space_before=0, first=False, align=None, line=1.35, italic=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    p.line_spacing = line
    if align is not None:
        p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = font
    return p


def rich(tf, parts, size=14, space_after=6, first=False, line=1.35, font=SANS):
    """A paragraph made of several differently-styled runs."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after = Pt(space_after)
    p.line_spacing = line
    for text, bold, color in parts:
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = font
    return p


def bullets(tf, items, size=13.5, color=INK2, space_after=9, bullet="—  "):
    for it in items:
        if isinstance(it, tuple):
            rich(tf, [(bullet, False, INK3), (it[0], True, INK), (it[1], False, color)],
                 size=size, space_after=space_after)
        else:
            para(tf, bullet + it, size=size, color=color, space_after=space_after)


def eyebrow(s, text, y=None):
    tf = textbox(s, M, y or Inches(0.52), CONTENT_W, Inches(0.26))
    para(tf, text.upper(), size=10.5, color=INK3, bold=True, first=True, space_after=0)


def title(s, text, y=Inches(0.84), size=32, w=None, color=INK):
    tf = textbox(s, M, y, w or CONTENT_W, Inches(1.0))
    para(tf, text, size=size, color=color, bold=False, font=SERIF, first=True,
         space_after=0, line=1.12)


def subtitle(s, text, y=Inches(1.66), w=None, size=14.5, color=INK2):
    tf = textbox(s, M, y, w or Inches(9.6), Inches(0.7))
    para(tf, text, size=size, color=color, first=True, space_after=0, line=1.4)


def footer(s, label=""):
    slide_no["n"] += 1
    tf = textbox(s, M, H - Inches(0.52), CONTENT_W, Inches(0.26))
    rich(tf, [("GCC Atlas", True, INK3),
              ("   ·   " + label if label else "", False, INK3)],
         size=9.5, first=True, space_after=0)
    tf2 = textbox(s, W - M - Inches(1.0), H - Inches(0.52), Inches(1.0), Inches(0.26),
                  align=PP_ALIGN.RIGHT)
    para(tf2, str(slide_no["n"]), size=9.5, color=INK3, first=True, space_after=0)


def rule(s, y, x=None, w=None, color=LINE):
    ln = s.shapes.add_shape(1, x or M, y, w or CONTENT_W, Pt(0.9))
    ln.fill.solid()
    ln.fill.fore_color.rgb = color
    ln.line.fill.background()
    ln.shadow.inherit = False
    return ln


def card(s, x, y, w, h, fill=WHITE, border=LINE):
    c = s.shapes.add_shape(1, x, y, w, h)
    c.fill.solid()
    c.fill.fore_color.rgb = fill
    c.line.color.rgb = border
    c.line.width = Pt(0.9)
    c.shadow.inherit = False
    return c


def chip(s, x, y, text, fg, bg, w=None):
    cw = w or Inches(0.05 * len(text) + 0.34)
    c = s.shapes.add_shape(1, x, y, cw, Inches(0.26))
    c.fill.solid()
    c.fill.fore_color.rgb = bg
    c.line.fill.background()
    c.shadow.inherit = False
    tf = c.text_frame
    tf.margin_left = tf.margin_right = Inches(0.07)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size = Pt(9.5)
    r.font.bold = True
    r.font.color.rgb = fg
    r.font.name = SANS
    return c


def crop_shot(name, out, top=0, height=None, width=None):
    """Crop a full-page screenshot down to the interesting part."""
    src = os.path.join(SHOTS, name)
    im = Image.open(src)
    w, h = im.size
    width = width or w
    height = height or min(h - top, int(width * 0.62))
    box = (0, top, width, min(top + height, h))
    im.crop(box).save(os.path.join(CROPS, out))
    return os.path.join(CROPS, out)


def shot_slide(s, img_path, x, y, width):
    pic = s.shapes.add_picture(img_path, x, y, width=width)
    # thin border so the screenshot reads as a framed object
    b = s.shapes.add_shape(1, x, y, pic.width, pic.height)
    b.fill.background()
    b.line.color.rgb = LINE
    b.line.width = Pt(1.1)
    b.shadow.inherit = False
    return pic


# ================================================================ SLIDE 1 ====
s = add_slide(WHITE)
bar = s.shapes.add_shape(1, 0, 0, Inches(0.13), H)
bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()
bar.shadow.inherit = False

tf = textbox(s, Inches(1.25), Inches(2.0), Inches(10.5), Inches(1.0))
para(tf, "GCC Atlas", size=54, color=INK, font=SERIF, first=True, space_after=0, line=1.0)

tf = textbox(s, Inches(1.25), Inches(3.0), Inches(9.4), Inches(1.1))
para(tf, "A source-backed view of India's GCC activity, built for the people "
         "who hire into these centres.",
     size=19, color=INK2, first=True, space_after=0, line=1.36)

rule(s, Inches(4.35), x=Inches(1.25), w=Inches(2.0), color=ACCENT)

tf = textbox(s, Inches(1.25), Inches(4.65), Inches(9.4), Inches(1.2))
para(tf, "Product Intern Assignment  ·  Flexiple", size=13, color=INK3, bold=True, first=True)
para(tf, "Working prototype + research dataset of 20 publicly reported GCC events", size=13, color=INK3)
para(tf, "Research date 10 October 2026", size=13, color=INK3)
footer(s, "Title")

# ================================================================ SLIDE 2 ====
s = add_slide()
eyebrow(s, "The short version")
title(s, "What I built, and why")
rule(s, Inches(1.56))

cols = [
    ("The problem", "Basic facts about India's GCC market are contested. Three "
     "credible sources give three different totals, and a single announcement is "
     "reported with different dates and different headcounts.", ACCENT),
    ("The user", "Recruitment and staffing firms selling into GCCs — the segment "
     "Flexiple itself sits in. They need to know who is building, and to be able "
     "to check it before they act on it.", ACCENT),
    ("The product", "A curated, source-backed record of GCC activity. Every claim "
     "carries its source, a confidence label, and a clear line between reported "
     "fact and my own interpretation.", ACCENT),
    ("The honest part", "20 records, not 2,000. No live data, no prediction, no "
     "user research yet. The product says so on its own pages rather than hiding it.", RED),
]
x = M
cw = (CONTENT_W - Inches(0.42)) / 4
for i, (h, body, col) in enumerate(cols):
    card(s, x, Inches(1.95), cw, Inches(3.5))
    acc = s.shapes.add_shape(1, x, Inches(1.95), cw, Pt(2.6))
    acc.fill.solid(); acc.fill.fore_color.rgb = col; acc.line.fill.background()
    acc.shadow.inherit = False
    tf = textbox(s, x + Inches(0.2), Inches(2.22), cw - Inches(0.4), Inches(3.0))
    para(tf, h, size=14.5, color=INK, bold=True, first=True, space_after=9)
    para(tf, body, size=12, color=INK2, line=1.42)
    x += cw + Inches(0.14)

tf = textbox(s, M, Inches(5.72), CONTENT_W, Inches(0.9))
rich(tf, [("One line: ", True, INK),
          ("GCC Atlas does not try to know more than the public record. It tries to "
           "be the only place where the public record is organised, dated, sourced "
           "and honest about its own gaps.", False, INK2)],
     size=13.5, first=True)
footer(s, "Summary")

# ================================================================ SLIDE 3 ====
s = add_slide()
eyebrow(s, "Problem")
title(s, "Who is in this ecosystem, and who feels the gap")
subtitle(s, "A company setting up an India GCC touches most of these groups. Each one "
            "needs a different slice of the same missing information.")
rule(s, Inches(2.32))

groups = [
    ("Global company setting up", "Which city, what it costs, who to trust", False),
    ("Advisors & consultants", "Who is in market, and when", False),
    ("Real estate firms", "Who needs space, and how much", False),
    ("Law firms", "Entity structure, compliance", False),
    ("State governments", "Who is shopping, and what they want", False),
    ("Recruitment & staffing firms", "Who is hiring, for what, starting when", True),
    ("EOR providers", "Who needs an employer of record first", False),
]
x, y = M, Inches(2.62)
cw2 = (CONTENT_W - Inches(0.36)) / 4
ch = Inches(1.12)
for i, (name, need, focus) in enumerate(groups):
    if i == 4:
        x, y = M, y + ch + Inches(0.18)
    card(s, x, y, cw2, ch,
         fill=ACCENT_BG if focus else WHITE,
         border=ACCENT if focus else LINE)
    tf = textbox(s, x + Inches(0.17), y + Inches(0.17), cw2 - Inches(0.34), ch - Inches(0.3))
    para(tf, name, size=12.5, color=ACCENT if focus else INK, bold=True, first=True, space_after=5)
    para(tf, need, size=11, color=INK2, line=1.35)
    x += cw2 + Inches(0.12)

tf = textbox(s, M, Inches(5.95), CONTENT_W, Inches(0.8))
rich(tf, [("I chose the highlighted group. ", True, ACCENT),
          ("They feel this weekly rather than once a decade, they already pay for "
           "market data, and they are reachable. They are also the segment Flexiple "
           "operates in — it runs a GCC recruitment and staffing business — so the "
           "product has a user inside the building.", False, INK2)],
     size=13, first=True)
footer(s, "Problem · Stakeholders")

# ================================================================ SLIDE 4 ====
s = add_slide()
eyebrow(s, "Evidence")
title(s, "One basic question, three credible answers")
subtitle(s, "\"How many GCCs are there in India?\" is the simplest question in this "
            "market. Here is what three reasonable sources say.")
rule(s, Inches(2.32))

nums = [
    ("2,000+", "Flexiple assignment brief", "The figure given in the brief for this assignment.", ACCENT),
    ("1,800+", "Flexiple's own website", "Stated on Flexiple's city-wise GCC directory page, alongside ~3,000 delivery units.", ACCENT),
    ("~1,700", "Vestian Research", "Reported via IANS, July 2025.", ACCENT),
]
x = M
cw3 = (CONTENT_W - Inches(0.3)) / 3
for n, src, note, col in nums:
    card(s, x, Inches(2.62), cw3, Inches(1.85))
    tf = textbox(s, x + Inches(0.26), Inches(2.84), cw3 - Inches(0.5), Inches(1.5))
    para(tf, n, size=38, color=col, bold=True, font=SERIF, first=True, space_after=4, line=1.0)
    para(tf, src, size=12.5, color=INK, bold=True, space_after=5)
    para(tf, note, size=10.5, color=INK3, line=1.35)
    x += cw3 + Inches(0.15)

tf = textbox(s, M, Inches(4.75), CONTENT_W, Inches(1.5))
para(tf, "None of these sources is wrong. They use different definitions, different "
         "dates and different counting methods — and none of them shows its working. "
         "That is the problem in one slide: not that information is missing, but that "
         "it cannot be checked.",
     size=14, color=INK, first=True, space_after=10, line=1.45)
rich(tf, [("Why this matters commercially:  ", True, ACCENT),
          ("a staffing firm sizing this market picks a number and builds a plan on it. "
           "A 300-centre spread between credible sources is the difference between a "
           "sound forecast and a guess.", False, INK2)],
     size=12.5)
footer(s, "Evidence · The counting problem")

# ================================================================ SLIDE 5 ====
s = add_slide()
eyebrow(s, "Evidence")
title(s, "What I hit while building the dataset")
subtitle(s, "I set out to verify 20 GCC announcements. The obstacles I ran into are "
            "themselves evidence of the problem — these are my own findings, not "
            "quoted research.")
rule(s, Inches(2.4))

findings = [
    ("Six different names for the same thing",
     "GCC, Global Capability Centre, Global Value Centre, Global Innovation Centre, "
     "capability center, Global Business Solutions — all in my 20 records. Any keyword "
     "search misses centres purely because of vocabulary."),
    ("The same event, different dates",
     "ICE's Pune lease: one source dates it 29 Dec 2025 (registration), another "
     "Feb 2026 (reporting). Both are probably right about different events."),
    ("The same event, different headcounts",
     "Eli Lilly's Hyderabad centre: 1,000+ at announcement, 1,500 at inauguration. "
     "HCA Healthcare: ~1,100 in one report, ~1,200 in another."),
    ("Headline numbers that mean something else",
     "JPMorgan Chase's Mumbai campus is widely reported as \"30,000 employees\". "
     "That is building capacity for a 2029 campus, and the reporting says "
     "\"consolidating\" existing India operations. It is not 30,000 new jobs."),
    ("Most coverage has no primary source",
     "Only 8 of my 20 records cite a company press release. Much of what circulates "
     "is LinkedIn posts and \"Top 10 GCCs\" roundups with nothing underneath them."),
]
y = Inches(2.74)
for i, (head, body) in enumerate(findings):
    num = s.shapes.add_shape(1, M, y + Inches(0.02), Inches(0.3), Inches(0.3))
    num.fill.solid(); num.fill.fore_color.rgb = ACCENT; num.line.fill.background()
    num.shadow.inherit = False
    ntf = num.text_frame
    ntf.margin_left = ntf.margin_right = ntf.margin_top = ntf.margin_bottom = 0
    ntf.vertical_anchor = MSO_ANCHOR.MIDDLE
    np_ = ntf.paragraphs[0]; np_.alignment = PP_ALIGN.CENTER
    nr = np_.add_run(); nr.text = str(i + 1)
    nr.font.size = Pt(11); nr.font.bold = True; nr.font.color.rgb = WHITE; nr.font.name = SANS

    tf = textbox(s, M + Inches(0.46), y, CONTENT_W - Inches(0.46), Inches(0.8))
    para(tf, head, size=13, color=INK, bold=True, first=True, space_after=3)
    para(tf, body, size=11.5, color=INK2, line=1.35)
    y += Inches(0.84)
footer(s, "Evidence · Research findings")

# ================================================================ SLIDE 6 ====
s = add_slide()
eyebrow(s, "Current alternatives")
title(s, "How people solve this today — and where it breaks")
rule(s, Inches(1.56))

rows = [
    ("Industry reports (Nasscom, Vestian, consultancies)",
     "Credible, broad market picture", "Annual, aggregate, not per-company, method rarely shown", "Verified"),
    ("Directory and listicle pages, including Flexiple's own",
     "Fast, free, good coverage of names", "A list of names, no dates, no sources per entry, no confidence", "Verified"),
    ("Google News / general search",
     "Current and free", "Terminology problem means you miss centres; no way to tell a press release from a blog", "Verified"),
    ("LinkedIn posts and roundups",
     "Early and specific", "Usually unsourced; figures are repeated without checking", "Verified"),
    ("Personal networks and referrals",
     "Genuinely the best current tool", "Does not scale, cannot be audited, excludes new entrants", "Hypothesis"),
    ("Sales intelligence tools (Apollo, ZoomInfo)",
     "Strong contact data", "No concept of \"GCC\"; cannot answer who is building a centre", "Hypothesis"),
]
hdr_y = Inches(1.9)
widths = [Inches(3.7), Inches(2.9), Inches(4.0), Inches(0.82)]
labels = ["Alternative", "What works", "Where it breaks", ""]
x = M
for lw, lab in zip(widths, labels):
    tf = textbox(s, x, hdr_y, lw, Inches(0.25))
    para(tf, lab.upper(), size=9.5, color=INK3, bold=True, first=True, space_after=0)
    x += lw + Inches(0.12)
rule(s, hdr_y + Inches(0.3))

y = hdr_y + Inches(0.44)
for alt, good, bad, status in rows:
    x = M
    tf = textbox(s, x, y, widths[0], Inches(0.6))
    para(tf, alt, size=11.5, color=INK, bold=True, first=True, space_after=0, line=1.3)
    x += widths[0] + Inches(0.12)
    tf = textbox(s, x, y, widths[1], Inches(0.6))
    para(tf, good, size=11, color=INK2, first=True, space_after=0, line=1.3)
    x += widths[1] + Inches(0.12)
    tf = textbox(s, x, y, widths[2], Inches(0.6))
    para(tf, bad, size=11, color=INK2, first=True, space_after=0, line=1.3)
    x += widths[2] + Inches(0.12)
    if status == "Verified":
        chip(s, x, y - Inches(0.02), "Checked", GREEN, GREEN_BG, w=Inches(0.82))
    else:
        chip(s, x, y - Inches(0.02), "Assumed", AMBER, AMBER_BG, w=Inches(0.82))
    y += Inches(0.63)

tf = textbox(s, M, Inches(6.22), CONTENT_W, Inches(0.6))
rich(tf, [("Reading the labels: ", True, INK),
          ("\"Checked\" means I used the alternative myself while researching this "
           "dataset. \"Assumed\" means it is a reasonable belief I have not tested — "
           "I have not interviewed anyone who does this job.", False, INK2)],
     size=11.5, first=True)
footer(s, "Current alternatives")

# ================================================================ SLIDE 7 ====
s = add_slide()
eyebrow(s, "Opportunities")
title(s, "Four things I could have built. I built the third.")
rule(s, Inches(1.56))

opts = [
    ("Provider track records",
     "\"Which advisor or EOR actually delivers?\"",
     "Rejected", "The most-felt pain, but the data is private outcomes. Building it "
     "honestly is impossible; building it dishonestly means inventing performance data.", RED, RED_BG),
    ("Setup decision copilot",
     "\"Which city, what entity, what cost?\"",
     "Rejected", "Very few buyers, decade-long cycles, and it depends on verified cost "
     "and incentive figures I could not source reliably in the time available.", RED, RED_BG),
    ("Source-backed activity record",
     "\"Who is building, and can I check it?\"",
     "Built", "Entirely buildable from public evidence. Serves a user who feels the "
     "need weekly. Honest about what it does not know.", GREEN, GREEN_BG),
    ("Expansion prediction",
     "\"Who will set up a GCC next?\"",
     "Rejected", "I could not find evidence that public data gives reliable early "
     "warning. Claiming prediction without testing it would be the dishonest option.", RED, RED_BG),
]
x = M
cw4 = (CONTENT_W - Inches(0.45)) / 4
for name, q, verdict, why, col, colbg in opts:
    h = Inches(3.9)
    card(s, x, Inches(1.95), cw4, h,
         fill=WHITE if verdict == "Rejected" else GREEN_BG,
         border=GREEN if verdict == "Built" else LINE)
    tf = textbox(s, x + Inches(0.2), Inches(2.18), cw4 - Inches(0.4), Inches(3.4))
    para(tf, name, size=13.5, color=INK, bold=True, first=True, space_after=6, line=1.25)
    para(tf, q, size=11, color=INK3, italic=True, space_after=10, line=1.3)
    para(tf, why, size=11, color=INK2, line=1.4)
    chip(s, x + Inches(0.2), Inches(5.42), verdict, col, colbg,
         w=Inches(0.78) if verdict == "Built" else Inches(0.92))
    x += cw4 + Inches(0.15)

tf = textbox(s, M, Inches(6.08), CONTENT_W, Inches(0.8))
rich(tf, [("The deciding question was the same each time: ", True, INK),
          ("can this be built honestly from evidence I can actually obtain? Only one "
           "of the four could. Three of these would have made a more impressive "
           "slide and a less defensible product.", False, INK2)],
     size=13, first=True)
footer(s, "Opportunities · Prioritisation")

# ================================================================ SLIDE 8 ====
s = add_slide()
eyebrow(s, "Target user")
title(s, "The first user, and the job to be done")
rule(s, Inches(1.56))

card(s, M, Inches(1.95), Inches(5.9), Inches(3.5), fill=ACCENT_BG, border=ACCENT)
tf = textbox(s, M + Inches(0.32), Inches(2.22), Inches(5.3), Inches(3.0))
para(tf, "Business development lead at a recruitment or staffing firm",
     size=17, color=ACCENT, bold=True, font=SERIF, first=True, space_after=11, line=1.2)
para(tf, "Their job is to find accounts worth approaching. Today that means reading "
         "news, watching LinkedIn, and relying on who they happen to know.",
     size=12.5, color=INK2, space_after=12, line=1.42)
rich(tf, [("The job to be done:  ", True, ACCENT),
          ("\"Show me which companies are building or growing a capability centre in "
           "India, tell me when, and let me check it before I put my name to it.\"",
           False, INK)], size=12.5)

x2 = M + Inches(6.2)
reasons = [
    ("Felt weekly, not once a decade", "Frequency is what makes a tool stick."),
    ("Already pays for market data", "No new budget line to create — a better one to argue for."),
    ("Reachable", "Thousands of firms, no procurement committee."),
    ("Serveable with honest data", "Public evidence genuinely answers their question."),
    ("Flexiple is this user", "It runs a GCC recruitment and staffing business — it could use this itself."),
]
y = Inches(1.95)
for h, b in reasons:
    card(s, x2, y, CONTENT_W - Inches(6.2), Inches(0.63))
    tf = textbox(s, x2 + Inches(0.18), y + Inches(0.1), CONTENT_W - Inches(6.6), Inches(0.5))
    rich(tf, [(h + "  ", True, INK), (b, False, INK2)], size=11.5, first=True, line=1.25)
    y += Inches(0.72)
footer(s, "Target user")

# ================================================================ SLIDE 9 ====
s = add_slide()
eyebrow(s, "Solution")
title(s, "GCC Atlas")
subtitle(s, "A curated record of publicly reported GCC activity in India where every "
            "claim can be traced, questioned and checked.", size=15.5)
rule(s, Inches(2.4))

princ = [
    ("Every claim carries its source",
     "47 source links across 20 records, each tagged as a company release, trade "
     "press or news, with its date. All clickable."),
    ("Facts and interpretation never blur",
     "Each record shows what the sources state, then — separately and visibly "
     "labelled — what I read into it. The two are styled differently so you "
     "cannot mistake one for the other."),
    ("Confidence is about evidence, not importance",
     "A small expansion with a company press release outranks a huge announcement "
     "carried only by a blog. The reason for each rating is written out."),
    ("It refuses rather than guesses",
     "Ask it about salaries, costs, incentives or the future and it names what it "
     "does not hold instead of producing something plausible."),
]
x = M
cw5 = (CONTENT_W - Inches(0.45)) / 4
for h, b in princ:
    card(s, x, Inches(2.72), cw5, Inches(2.5))
    acc = s.shapes.add_shape(1, x, Inches(2.72), cw5, Pt(2.6))
    acc.fill.solid(); acc.fill.fore_color.rgb = ACCENT; acc.line.fill.background()
    acc.shadow.inherit = False
    tf = textbox(s, x + Inches(0.2), Inches(2.98), cw5 - Inches(0.4), Inches(2.1))
    para(tf, h, size=13, color=INK, bold=True, first=True, space_after=8, line=1.25)
    para(tf, b, size=11, color=INK2, line=1.42)
    x += cw5 + Inches(0.15)

tf = textbox(s, M, Inches(5.52), CONTENT_W, Inches(1.2))
para(tf, "What it is not", size=11, color=RED, bold=True, first=True, space_after=6)
para(tf, "Not live data — the dataset is a fixed snapshot in a file.   ·   "
         "Not a complete register — 20 records, not 2,000.   ·   "
         "Not a prediction tool.   ·   Not validated by any real user yet.",
     size=12, color=INK2, line=1.5)
footer(s, "Solution")

# ============================================================ SLIDES 10-13 ===
shots = [
    ("01-overview.png", "The overview", 0, 700,
     "Every number here is counted from the 20 records at page load — nothing is "
     "hardcoded, so the figures cannot drift from the data. The banner states the "
     "dataset's size and limits before showing a single chart."),
    ("05-company-detail.png", "Evidence, separated", 150, 700,
     "The blue-marked block is what the sources say. The purple block below it is "
     "labelled \"our interpretation — not from the sources\". Headcount claims carry "
     "their basis. Every source is one click away."),
    ("07-ask-refusal.png", "Ask: refusing well", 110, 660,
     "Asked about salaries, it names exactly what it does not hold and why, then "
     "points to the method page. It cites nothing, because it has nothing. The panel "
     "below states plainly that this is keyword search, not AI."),
    ("10-verify.png", "Verification built in", 90, 700,
     "All 47 source links in one checklist, so the research can be audited link by "
     "link. This page exists because the dataset was assembled from search results "
     "and the links still need first-hand confirmation."),
]
for fname, head, top, height, caption in shots:
    s = add_slide()
    eyebrow(s, "Prototype")
    title(s, head, size=27)
    tf = textbox(s, M, Inches(1.5), Inches(11.6), Inches(0.62))
    para(tf, caption, size=12.5, color=INK2, first=True, space_after=0, line=1.4)
    crop = crop_shot(fname, "crop-" + fname, top=top, height=height)
    img = Image.open(crop)
    max_w, max_h = Inches(10.9), Inches(4.72)
    ratio = img.size[1] / img.size[0]
    w_ = max_w
    h_ = Emu(int(w_ * ratio))
    if h_ > max_h:
        h_ = max_h
        w_ = Emu(int(h_ / ratio))
    shot_slide(s, crop, int((W - w_) / 2), Inches(2.3), w_)
    footer(s, "Prototype walkthrough")

# =============================================================== SLIDE 14 ====
s = add_slide()
eyebrow(s, "Scope")
title(s, "Why these features, and what I deliberately cut")
rule(s, Inches(1.56))

tf = textbox(s, M, Inches(1.9), Inches(6.0), Inches(0.3))
para(tf, "BUILT — because the journey breaks without it", size=10, color=GREEN, bold=True, first=True, space_after=0)
built = [
    ("Overview", "orientation: what is in here, and how much"),
    ("Activity table + filters", "the actual working surface"),
    ("Company detail", "where trust is won or lost"),
    ("Ask", "a faster route in for a specific question"),
    ("City breakdown", "pattern view, counted only from these records"),
    ("Sources & Method", "a tool that won't explain itself can't be trusted"),
    ("Verify checklist", "makes the research auditable"),
]
y = Inches(2.22)
for n, why in built:
    card(s, M, y, Inches(6.0), Inches(0.52), fill=GREEN_BG, border=GREEN_BG)
    tf = textbox(s, M + Inches(0.16), y + Inches(0.1), Inches(5.7), Inches(0.4))
    rich(tf, [(n + " — ", True, INK), (why, False, INK2)], size=11, first=True, line=1.2)
    y += Inches(0.6)

x2 = M + Inches(6.4)
tf = textbox(s, x2, Inches(1.9), Inches(5.3), Inches(0.3))
para(tf, "CUT — and why", size=10, color=RED, bold=True, first=True, space_after=0)
cut = [
    ("State incentive comparison", "the most error-prone data in this market; wrong is worse than absent"),
    ("Office cost benchmarks", "could not source reliably enough to publish"),
    ("Contact data", "a data-protection question I should not answer casually"),
    ("Live ingestion", "real engineering; not demonstrable in the time"),
    ("Login and saved accounts", "nothing here is private"),
    ("A database", "20 records in a file is the right answer; a database adds steps, not ability"),
    ("A framework", "no build step means it opens by double-clicking"),
]
y = Inches(2.22)
for n, why in cut:
    card(s, x2, y, Inches(5.3), Inches(0.52), fill=WHITE)
    tf = textbox(s, x2 + Inches(0.16), y + Inches(0.1), Inches(5.0), Inches(0.4))
    rich(tf, [(n + " — ", True, INK), (why, False, INK2)], size=11, first=True, line=1.2)
    y += Inches(0.6)

tf = textbox(s, M, Inches(6.3), CONTENT_W, Inches(0.5))
rich(tf, [("The build: ", True, ACCENT),
          ("plain HTML, CSS and JavaScript. No framework, no server, no database, no "
           "API key, no paid service. It opens by double-clicking a file and works "
           "offline. I can explain every line of it.", False, INK2)],
     size=12.5, first=True)
footer(s, "Scope & trade-offs")

# =============================================================== SLIDE 15 ====
s = add_slide()
eyebrow(s, "Positioning")
title(s, "Why use this instead of today's alternatives", size=30)
rule(s, Inches(1.56))

comp = [
    ("Industry reports", "Market-level picture, annually",
     "Per-company and dated, so it supports an action this week"),
    ("Directory pages\n(including Flexiple's)", "A list of company names",
     "Each entry carries its source, date, confidence and caveats — a list becomes checkable evidence"),
    ("Google News search", "Everything, unranked and unlabelled",
     "Curated across six naming conventions, with weak sources excluded on purpose"),
    ("LinkedIn roundups", "Fast, specific, usually unsourced",
     "Records without a named publication behind them were left out"),
    ("A general AI chatbot", "Fluent answers, unverifiable",
     "Answers only from 20 cited records, and refuses when it has nothing"),
]
hdr_y = Inches(1.95)
w1, w2, w3 = Inches(2.9), Inches(3.5), Inches(5.2)
x = M
for lw, lab in zip([w1, w2, w3], ["They use today", "What it gives them", "What GCC Atlas adds"]):
    tf = textbox(s, x, hdr_y, lw, Inches(0.25))
    para(tf, lab.upper(), size=9.5, color=INK3, bold=True, first=True, space_after=0)
    x += lw + Inches(0.2)
rule(s, hdr_y + Inches(0.3))

y = hdr_y + Inches(0.46)
for a, b, c in comp:
    tf = textbox(s, M, y, w1, Inches(0.7))
    para(tf, a, size=11.5, color=INK, bold=True, first=True, space_after=0, line=1.25)
    tf = textbox(s, M + w1 + Inches(0.2), y, w2, Inches(0.7))
    para(tf, b, size=11, color=INK3, first=True, space_after=0, line=1.3)
    tf = textbox(s, M + w1 + w2 + Inches(0.4), y, w3, Inches(0.7))
    para(tf, c, size=11, color=INK, first=True, space_after=0, line=1.3)
    y += Inches(0.76)

card(s, M, Inches(6.12), CONTENT_W, Inches(0.72), fill=ACCENT_BG, border=ACCENT)
tf = textbox(s, M + Inches(0.24), Inches(6.26), CONTENT_W - Inches(0.48), Inches(0.5))
rich(tf, [("A note on Flexiple's own pages: ", True, ACCENT),
          ("these are good directories and they rank well. GCC Atlas is not a "
           "replacement — it is the research layer underneath, where each entry "
           "earns its place by citing something.", False, INK2)],
     size=11.5, first=True, line=1.3)
footer(s, "Positioning")

# =============================================================== SLIDE 16 ====
s = add_slide()
eyebrow(s, "Honesty")
title(s, "What is not true about this yet")
subtitle(s, "Stated here rather than discovered later. The product repeats all of this "
            "on its own Sources & Method page.")
rule(s, Inches(2.32))

limits = [
    ("The source links are not yet confirmed first-hand",
     "I identified them through search. The build environment blocked opening pages "
     "directly, so each link still needs to be read and confirmed. The app ships a "
     "Verify checklist for exactly this, and nothing should be presented as fact "
     "until that is done.", RED),
    ("No user research exists",
     "I have not interviewed a single recruitment or staffing professional. Every "
     "statement about what they need is my reasoning, not a finding. The alternatives "
     "slide marks which claims I tested myself and which I assumed.", RED),
    ("20 records is a sample, not a market",
     "Coverage is skewed toward cities and states with active press communications. "
     "A count here measures reporting volume, not real activity.", AMBER),
    ("It cannot predict, and does not try",
     "I looked for evidence that public signals give reliable early warning of GCC "
     "setups. I did not find enough to support that claim, so the product does not "
     "make it.", AMBER),
    ("\"Ask\" is keyword search, not AI",
     "It scores records by word overlap and detects a few patterns. The app says so "
     "on the page itself. No model, no API key, no cost.", AMBER),
]
y = Inches(2.62)
for head, body, col in limits:
    bar_ = s.shapes.add_shape(1, M, y, Pt(2.6), Inches(0.72))
    bar_.fill.solid(); bar_.fill.fore_color.rgb = col; bar_.line.fill.background()
    bar_.shadow.inherit = False
    tf = textbox(s, M + Inches(0.22), y, CONTENT_W - Inches(0.22), Inches(0.72))
    para(tf, head, size=12.5, color=INK, bold=True, first=True, space_after=3)
    para(tf, body, size=11, color=INK2, line=1.33)
    y += Inches(0.86)
footer(s, "Limitations & assumptions")

# =============================================================== SLIDE 17 ====
s = add_slide()
eyebrow(s, "Next")
title(s, "What I would do next, in order")
rule(s, Inches(1.56))

steps = [
    ("1", "Confirm all 47 sources", "Half a day",
     "Open every link, confirm it supports the record, log who checked it and when. "
     "Nothing else matters until this is done."),
    ("2", "Talk to eight people who do this job", "One week",
     "Recruitment and staffing BD leads. One question above all: what do you use "
     "today, and what would make you stop using it? This is the assumption the whole "
     "product rests on."),
    ("3", "Test whether state MoUs are an early signal", "One week",
     "One record here shows roughly 17 months between a state MoU and an opening. "
     "If that holds across 20 more cases, it is a genuinely earlier signal than "
     "anything in the dataset today."),
    ("4", "Widen coverage with a repeatable method", "Two weeks",
     "Move from manual searching to a defined collection process, and record why "
     "each candidate was included or excluded."),
    ("5", "Add company registry data", "Later",
     "Tie each centre to its Indian legal entity, so the same company is never "
     "counted twice under two names."),
]
y = Inches(1.95)
for num, head, effort, body in steps:
    circ = s.shapes.add_shape(9, M, y + Inches(0.06), Inches(0.36), Inches(0.36))  # 9 = oval
    circ.fill.solid(); circ.fill.fore_color.rgb = ACCENT; circ.line.fill.background()
    circ.shadow.inherit = False
    ctf = circ.text_frame
    ctf.margin_left = ctf.margin_right = ctf.margin_top = ctf.margin_bottom = 0
    ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
    cp = ctf.paragraphs[0]; cp.alignment = PP_ALIGN.CENTER
    cr = cp.add_run(); cr.text = num
    cr.font.size = Pt(12); cr.font.bold = True; cr.font.color.rgb = WHITE; cr.font.name = SANS

    tf = textbox(s, M + Inches(0.54), y, Inches(9.2), Inches(0.8))
    para(tf, head, size=13.5, color=INK, bold=True, first=True, space_after=3)
    para(tf, body, size=11.5, color=INK2, line=1.35)

    chip(s, W - M - Inches(1.3), y + Inches(0.05), effort, INK3, RGBColor(0xF4, 0xF2, 0xEE),
         w=Inches(1.3))
    y += Inches(0.94)
footer(s, "Next steps")

# =============================================================== SLIDE 18 ====
s = add_slide(WHITE)
bar = s.shapes.add_shape(1, 0, 0, Inches(0.13), H)
bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()
bar.shadow.inherit = False

tf = textbox(s, Inches(1.25), Inches(1.85), Inches(10.4), Inches(1.0))
para(tf, "The argument in one line", size=15, color=INK3, bold=True, first=True, space_after=16)
para(tf, "This market does not lack information.\nIt lacks information you can check.",
     size=30, color=INK, font=SERIF, line=1.25, space_after=20)
rule(s, Inches(4.25), x=Inches(1.25), w=Inches(2.0), color=ACCENT)

tf = textbox(s, Inches(1.25), Inches(4.6), Inches(10.0), Inches(1.6))
para(tf, "GCC Atlas is 20 records, 47 sources and five cities. It is small on purpose. "
         "Every claim in it points at something you can open, every interpretation is "
         "labelled as mine, and everything it does not know is written down.",
     size=14.5, color=INK2, first=True, space_after=14, line=1.45)
rich(tf, [("Prototype: ", True, INK), ("open index.html — no install, no account, no internet needed.   ",
                                        False, INK2),
          ("Dataset: ", True, INK), ("data/records.js.", False, INK2)], size=12.5)
footer(s, "Close")

# ------------------------------------------------------------------- save ---
out = os.path.join(DECK_DIR, "GCC-Atlas-Flexiple-Assignment.pptx")
prs.save(out)
print("Deck written to: " + out)
print("Slides: %d" % len(prs.slides.__iter__.__self__._sldIdLst))
